"""Validate and synchronize uploaded spreadsheet data."""
import csv
import math
from pathlib import Path
from datetime import datetime
from zipfile import BadZipFile, ZipFile
import db

COLUMN_MAP = {
    "Request No":        "request_no",
    "Door Name":         "door_name",
    "Quotation No":      "quotation_no",
    "Vendor PO":         "vendor_po",
    "Brand":             "brand",
    "Request Date":      "request_date",
    "Installation Date": "installation_date",
    "Amount":            "amount",
    "Status":            "status",
    "SharePoint Folder": "sharepoint_folder",
}
MAX_SHEET_ROWS = 50_000
MAX_FIELD_LENGTHS = {
    "Request No": 128,
    "Door Name": 500,
    "Quotation No": 200,
    "Vendor PO": 200,
    "Brand": 200,
    "Request Date": 40,
    "Installation Date": 40,
    "Status": 80,
    "SharePoint Folder": 500,
}


def _validate_headers(headers: list[str]) -> list[str]:
    normalized = [str(value or "").strip() for value in headers]
    if not normalized or "Request No" not in normalized:
        raise ValueError("Sheet must contain a 'Request No' column.")
    if len(normalized) != len(set(normalized)):
        raise ValueError("Sheet contains duplicate column names.")
    if any(not header for header in normalized):
        raise ValueError("Sheet contains a blank column name.")
    return normalized


def _fmt_date(v):
    if v is None: return None
    if isinstance(v, datetime): return v.strftime("%Y-%m-%d")
    return str(v).strip()


def _load_xlsx(path: str) -> list[dict]:
    import openpyxl
    try:
        with ZipFile(path) as archive:
            if sum(entry.file_size for entry in archive.infolist()) > 100 * 1024 * 1024:
                raise ValueError("Expanded spreadsheet content exceeds the 100 MB limit.")
            if archive.testzip():
                raise ValueError("Spreadsheet archive is damaged.")
    except BadZipFile as exc:
        raise ValueError("The uploaded .xlsx file is not a valid spreadsheet.") from exc
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        ws = wb.active
        headers = _validate_headers([cell.value for cell in ws[1]])
        rows = []
        for raw in ws.iter_rows(min_row=2, values_only=True):
            if not raw or not raw[0]:
                continue
            rows.append(dict(zip(headers, raw)))
            if len(rows) > MAX_SHEET_ROWS:
                raise ValueError(f"Sheet cannot contain more than {MAX_SHEET_ROWS} rows.")
        return rows
    finally:
        wb.close()


def _load_xls(path: str) -> list[dict]:
    import xlrd
    wb = xlrd.open_workbook(path)
    ws = wb.sheet_by_index(0)
    headers = _validate_headers(
        [ws.cell_value(0, c) for c in range(ws.ncols)]
    )
    rows = []
    for r in range(1, ws.nrows):
        row = {headers[c]: ws.cell_value(r, c) for c in range(ws.ncols)}
        if not row.get("Request No"): continue
        rows.append(row)
        if len(rows) > MAX_SHEET_ROWS:
            raise ValueError(f"Sheet cannot contain more than {MAX_SHEET_ROWS} rows.")
    return rows


def _load_csv(path: str) -> list[dict]:
    with open(path, encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        if not reader.fieldnames:
            raise ValueError("CSV file is missing a header row.")
        headers = _validate_headers(reader.fieldnames)
        rows = []
        for row in reader:
            if not row.get("Request No"):
                continue
            rows.append({key: row.get(key) for key in headers})
            if len(rows) > MAX_SHEET_ROWS:
                raise ValueError(f"Sheet cannot contain more than {MAX_SHEET_ROWS} rows.")
        return rows


def load_sheet(path: str) -> list[dict]:
    suffix = Path(path).suffix.lower()
    if suffix == ".csv":
        rows = _load_csv(path)
    elif suffix == ".xls":
        rows = _load_xls(path)
    elif suffix == ".xlsx":
        rows = _load_xlsx(path)
    else:
        raise ValueError(f"Unsupported sheet type: {suffix or 'unknown'}")

    return rows


def run_sync(path: str) -> dict:
    """
    Compare sheet rows against tracker.
    Returns {"inserted": N, "updated": N, "skipped": N}
    """
    db.init_db()
    sheet_rows = load_sheet(path)
    if not sheet_rows:
        return {
            "inserted": 0,
            "updated": 0,
            "skipped": 0,
            "message": "Sheet is empty",
        }

    sheet_map = {}
    for row in sheet_rows:
        raw_request_no = row.get("Request No")
        if isinstance(raw_request_no, (int, float)) and not isinstance(raw_request_no, bool):
            if not math.isfinite(float(raw_request_no)):
                raise ValueError("Request numbers must be finite values.")
            req_no = (
                str(int(raw_request_no))
                if float(raw_request_no).is_integer()
                else str(raw_request_no)
            )
        else:
            req_no = str(raw_request_no or "").strip()
        if not req_no:
            continue
        if len(req_no) > MAX_FIELD_LENGTHS["Request No"]:
            raise ValueError("Request numbers cannot exceed 128 characters.")
        if req_no in sheet_map:
            raise ValueError(f"Duplicate request number in sheet: {req_no}")
        row["Request No"] = req_no
        for field, max_length in MAX_FIELD_LENGTHS.items():
            if field == "Request No":
                continue
            value = row.get(field)
            if value is None:
                continue
            if field in ("Request Date", "Installation Date"):
                value = _fmt_date(value)
            else:
                value = str(value).strip()
            if len(value) > max_length:
                raise ValueError(
                    f"{field} for request {req_no} exceeds {max_length} characters."
                )
            row[field] = value
        amount = row.get("Amount")
        if amount not in (None, ""):
            try:
                amount = float(str(amount).replace(",", "").replace("$", "").strip())
            except (TypeError, ValueError) as exc:
                raise ValueError(f"Invalid amount for request {req_no}.") from exc
            if not math.isfinite(amount) or amount < 0:
                raise ValueError(f"Invalid amount for request {req_no}.")
            row["Amount"] = amount
        sheet_map[req_no] = row

    conn = db.get_db()
    try:
        tracker_rows = db.get_all_tracker(conn)
        tracker_map  = {str(r["request_no"]): r for r in tracker_rows}

        to_insert = []
        to_update = []

        for req_no, sheet_row in sheet_map.items():
            if not req_no:
                continue

            if req_no not in tracker_map:
                to_insert.append(sheet_row)
            else:
                trk = tracker_map[req_no]
                trk_excel = {excel_col: trk.get(db_col)
                             for excel_col, db_col in COLUMN_MAP.items()}
                proposal = {
                    key: value for key, value in sheet_row.items()
                    if key in COLUMN_MAP
                    and trk_excel.get(key) in (None, "", "None", "none")
                    and value not in (None, "", "None", "none")
                }
                if proposal:
                    to_update.append((req_no, proposal))

        if not to_insert and not to_update:
            return {"inserted": 0, "updated": 0, "skipped": len(sheet_map),
                    "message": "Tracker is already up to date"}

        # Insert missing rows
        for row in to_insert:
            db_row = {COLUMN_MAP.get(k, k): v for k, v in row.items() if k in COLUMN_MAP}
            db_row["request_date"]      = _fmt_date(db_row.get("request_date"))
            db_row["installation_date"] = _fmt_date(db_row.get("installation_date"))
            db.upsert_request(conn, db_row, source="sync")
            if not db.insert_tracker_row(conn, row, source="sync"):
                raise ValueError(
                    f"Request {row.get('Request No')} was added concurrently. "
                    "Retry the upload."
                )

        # Fill in incomplete rows
        for req_no, proposal in to_update:
            db.update_tracker_row(conn, req_no, proposal, source="sync")

        conn.commit()

        return {
            "inserted": len(to_insert),
            "updated":  len(to_update),
            "skipped":  len(sheet_map) - len(to_insert) - len(to_update),
            "message":  f"Done — {len(to_insert)} added, {len(to_update)} updated",
        }
    finally:
        conn.close()
